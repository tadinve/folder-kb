import os
from typing import Optional

# PDF
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

# DOCX
try:
    import docx
except ImportError:
    docx = None

# XLSX
try:
    import openpyxl
except ImportError:
    openpyxl = None

def parse_txt(filepath: str) -> str:
    with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

def parse_pdf(filepath: str) -> str:
    if not PyPDF2:
        raise ImportError('PyPDF2 is not installed')
    text = ''
    with open(filepath, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        for page in reader.pages:
            text += page.extract_text() or ''
    return text

def parse_docx(filepath: str) -> str:
    if not docx:
        raise ImportError('python-docx is not installed')
    doc = docx.Document(filepath)
    return '\n'.join([p.text for p in doc.paragraphs])

def parse_xlsx(filepath: str) -> str:
    if not openpyxl:
        raise ImportError('openpyxl is not installed')
    wb = openpyxl.load_workbook(filepath, read_only=True)
    text = ''
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            text += '\t'.join([str(cell) if cell is not None else '' for cell in row]) + '\n'
    return text

def parse_file(filepath: str) -> Optional[str]:
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.txt':
        return parse_txt(filepath)
    elif ext == '.pdf':
        return parse_pdf(filepath)
    elif ext == '.docx':
        return parse_docx(filepath)
    elif ext == '.xlsx':
        return parse_xlsx(filepath)
    else:
        raise ValueError(f'Unsupported file type: {ext}')
